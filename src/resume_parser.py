import io
import os
from pathlib import Path
from typing import Optional, Union
import docx
from pypdf import PdfReader

from src.preprocessing import clean_text
from src.utils import get_logger

logger = get_logger("resume_parser")


def extract_text_from_pdf(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
    """
    Extract text content from a PDF document.

    Args:
        file_source: Filepath string/Path or in-memory BytesIO/bytes.

    Returns:
        str: Extracted and cleaned text.
    """
    try:
        if isinstance(file_source, (str, Path)):
            path = Path(file_source)
            if not path.exists():
                raise FileNotFoundError(f"PDF file not found: {path}")
            reader = PdfReader(str(path))
        elif isinstance(file_source, bytes):
            reader = PdfReader(io.BytesIO(file_source))
        elif isinstance(file_source, io.BytesIO):
            reader = PdfReader(file_source)
        elif hasattr(file_source, "read"):
            # Streamlit UploadedFile
            reader = PdfReader(io.BytesIO(file_source.read()))
        else:
            raise ValueError(f"Unsupported file source type: {type(file_source)}")

        extracted_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text)

        full_text = "\n\n".join(extracted_pages)
        return clean_text(full_text)
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {e}")
        raise ValueError(f"Failed to parse PDF document: {str(e)}") from e


def extract_text_from_docx(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
    """
    Extract text content from a Microsoft Word (.docx) document.

    Args:
        file_source: Filepath string/Path or in-memory BytesIO/bytes.

    Returns:
        str: Extracted and cleaned text.
    """
    try:
        if isinstance(file_source, (str, Path)):
            path = Path(file_source)
            if not path.exists():
                raise FileNotFoundError(f"DOCX file not found: {path}")
            doc = docx.Document(str(path))
        elif isinstance(file_source, bytes):
            doc = docx.Document(io.BytesIO(file_source))
        elif isinstance(file_source, io.BytesIO):
            doc = docx.Document(file_source)
        elif hasattr(file_source, "read"):
            doc = docx.Document(io.BytesIO(file_source.read()))
        else:
            raise ValueError(f"Unsupported file source type: {type(file_source)}")

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract table text if present
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n".join(paragraphs)
        return clean_text(full_text)
    except Exception as e:
        logger.error(f"Error extracting text from DOCX: {e}")
        raise ValueError(f"Failed to parse DOCX document: {str(e)}") from e


def extract_text_from_txt(file_source: Union[str, Path, io.BytesIO, bytes]) -> str:
    """
    Extract text content from a plain text (.txt) file with multi-encoding fallback.

    Args:
        file_source: Filepath string/Path or in-memory BytesIO/bytes.

    Returns:
        str: Extracted and cleaned text.
    """
    try:
        raw_bytes = None
        if isinstance(file_source, (str, Path)):
            path = Path(file_source)
            if not path.exists():
                raise FileNotFoundError(f"TXT file not found: {path}")
            with open(path, "rb") as f:
                raw_bytes = f.read()
        elif isinstance(file_source, bytes):
            raw_bytes = file_source
        elif isinstance(file_source, io.BytesIO):
            raw_bytes = file_source.getvalue()
        elif hasattr(file_source, "read"):
            raw_bytes = file_source.read()
            if isinstance(raw_bytes, str):
                return clean_text(raw_bytes)
        else:
            raise ValueError(f"Unsupported file source type: {type(file_source)}")

        # Try UTF-8 first, fallback to Latin-1
        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = raw_bytes.decode("latin-1", errors="replace")

        return clean_text(text)
    except Exception as e:
        logger.error(f"Error extracting text from TXT: {e}")
        raise ValueError(f"Failed to parse TXT document: {str(e)}") from e


def extract_resume_text(
    file_source: Union[str, Path, any],
    filename: Optional[str] = None,
) -> str:
    """
    Universal resume text extractor. Automatically detects format (PDF, DOCX, TXT)
    from filename extension or object attributes.

    Args:
        file_source: Filepath, Streamlit UploadedFile, or BytesIO.
        filename (Optional[str]): Explicit filename for extension detection.

    Returns:
        str: Clean extracted resume text.
    """
    fname = filename
    if fname is None:
        if isinstance(file_source, (str, Path)):
            fname = str(file_source)
        elif hasattr(file_source, "name"):
            fname = file_source.name

    if not fname:
        # Default fallback to TXT
        return extract_text_from_txt(file_source)

    fname_lower = fname.lower()
    if fname_lower.endswith(".pdf"):
        return extract_text_from_pdf(file_source)
    elif fname_lower.endswith(".docx"):
        return extract_text_from_docx(file_source)
    elif fname_lower.endswith(".txt"):
        return extract_text_from_txt(file_source)
    else:
        raise ValueError(f"Unsupported file extension: {fname}. Supported formats: PDF (.pdf), Word (.docx), Plain Text (.txt)")
